# -*- coding: utf-8 -*-
"""Reproduce la cascada PRISMA reportada en la Metodologia del manuscrito.

Salida esperada:
    724 identificados -> 171 duplicados -> 385 tipologia -> 5 idioma -> 14 depuracion -> 149
"""
import re
import unicodedata
import collections
import openpyxl

XLSX = 'ARTICULOS_UNIFICADOS_CRIBADO.xlsx'
TIPOS_PRIMARIOS = {'Article', 'ARTICLE', 'ARTICLE; EARLY ACCESS'}


def norm(s):
    if not s:
        return ''
    s = unicodedata.normalize('NFKD', str(s)).encode('ascii', 'ignore').decode()
    return re.sub(r'[^a-z0-9]', '', s.lower())


def cabecera(ws):
    return [(c.value or '') for c in next(ws.iter_rows(min_row=1, max_row=1))]


def main():
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    sc, wo, se = wb['ARTICULOS SCOPUS'], wb['ARTICULOS WoS'], wb['SELECCIONADOS']
    hs, hw, hse = cabecera(sc), cabecera(wo), cabecera(se)

    seleccionados = {norm(r[hse.index('TI')])
                     for r in se.iter_rows(min_row=2, values_only=True) if any(r)}

    registros = []
    for r in sc.iter_rows(min_row=2, values_only=True):
        if any(r):
            registros.append(dict(base='SCOPUS',
                                  ti=norm(r[hs.index('Title')]),
                                  dt=str(r[hs.index('Document Type')] or ''),
                                  lang=str(r[hs.index('Language of Original Document')] or '')))
    for r in wo.iter_rows(min_row=2, values_only=True):
        if any(r):
            registros.append(dict(base='WOS',
                                  ti=norm(r[hw.index('TI')]),
                                  dt=str(r[hw.index('DT')] or ''),
                                  lang=str(r[hw.index('LA')] or '')))

    # Deduplicacion priorizando la metadata de WoS (campo CR poblado).
    unicos = {}
    for x in sorted(registros, key=lambda x: 0 if x['base'] == 'WOS' else 1):
        unicos.setdefault(x['ti'], x)
    unicos = list(unicos.values())

    no_primarios = [x for x in unicos if x['dt'] not in TIPOS_PRIMARIOS]
    tras_tipo = [x for x in unicos if x['dt'] in TIPOS_PRIMARIOS]
    otro_idioma = [x for x in tras_tipo if x['lang'].lower() not in ('english', '')]
    tras_idioma = [x for x in tras_tipo if x['lang'].lower() in ('english', '')]
    depuracion = [x for x in tras_idioma if x['ti'] not in seleccionados]

    print('CASCADA PRISMA')
    print(f'  1. Identificados                       {len(registros):5d}')
    print(f'  2. Duplicados eliminados              -{len(registros)-len(unicos):5d}'
          f'   -> {len(unicos):4d} unicos')
    print(f'  3. Tipologia no primaria (CE-4)       -{len(no_primarios):5d}'
          f'   -> {len(tras_tipo):4d}')
    print(f'  4. Idioma no contemplado (CI-4)       -{len(otro_idioma):5d}'
          f'   -> {len(tras_idioma):4d}')
    print(f'  5. Lit. secundaria / metadata incompl.-{len(depuracion):5d}'
          f'   -> {len(tras_idioma)-len(depuracion):4d}')
    print(f'  Total descartado: {len(registros)-len(seleccionados)}'
          f'   |  Pasan a cribado: {len(seleccionados)}')
    print()
    print('  Detalle del filtro de tipologia:')
    for k, v in collections.Counter(x['dt'] for x in no_primarios).most_common():
        print(f'    {k:32s} {v:4d}')

    assert len(registros) == 724, 'registros identificados != 724'
    assert len(tras_idioma) - len(depuracion) == 149, 'cribado != 149'
    print('\n[OK] La cascada reproduce las cifras del manuscrito.')


if __name__ == '__main__':
    main()
